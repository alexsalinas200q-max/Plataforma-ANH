# apps/solicitudes/tests.py
#
# Cobertura de las dos reglas duras agregadas por validar_cupo.py:
# cupo mensual de 120L y unicidad de solicitud activa. No se testea
# acá el resto del ciclo de vida de Solicitud (aprobar/despachar/etc.),
# que no tiene suite propia todavía.

from datetime import date, datetime
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from catalogos.models import Departamento, Municipio, Provincia
from consumidores.models import ConsumidorPerfil
from estaciones.models import EstacionServicio
from solicitudes.models import Solicitud
from solicitudes.services.aprobar_solicitud import aprobar_solicitud
from solicitudes.services.validar_cupo import (
    calcular_cupo_mensual,
    obtener_solicitud_activa,
    verificar_puede_aprobar,
    verificar_puede_crear_solicitud,
)

User = get_user_model()


def _crear_consumidor(email, estado_identidad=None):
    user = User.objects.create_user(
        email=email,
        nombres="Test",
        apellido_paterno="Consumidor",
        tipo_usuario=User.TipoUsuario.CONS,
        password="testpass123",
    )
    perfil = ConsumidorPerfil.objects.create(
        user=user,
        fecha_nacimiento=date(1990, 1, 1),
    )
    if estado_identidad:
        perfil.estado_identidad = estado_identidad
        perfil.save(update_fields=["estado_identidad"])
    return perfil


def _crear_solicitud(
    consumidor,
    estado,
    litros_solicitados=50,
    litros_aprobados=None,
    fecha_aprobacion=None,
):
    # creado_por y declaracion_jurada_confirmada no son necesarios para
    # los tests que solo leen el estado/acumulado vía queryset, pero sí
    # los exige Solicitud.full_clean() — se setean siempre para que los
    # tests de integración (que sí llaman aprobar_solicitud(), y por lo
    # tanto full_clean()) no fallen por un fixture incompleto.
    return Solicitud.objects.create(
        consumidor=consumidor,
        creado_por=consumidor.user,
        tipo_combustible=Solicitud.TipoCombustible.GASOLINA,
        litros_solicitados=litros_solicitados,
        estado=estado,
        litros_aprobados=litros_aprobados,
        fecha_aprobacion=fecha_aprobacion,
        declaracion_jurada_confirmada=True,
        fecha_declaracion_jurada=timezone.now(),
    )


def _crear_cadena_geografica():
    departamento = Departamento.objects.create(nombre="La Paz", codigo="LP")
    provincia = Provincia.objects.create(
        departamento=departamento, nombre="Murillo", codigo="001",
    )
    municipio = Municipio.objects.create(
        provincia=provincia, nombre="La Paz", codigo="0101",
    )
    return departamento, provincia, municipio


def _crear_estacion(codigo="EST-TEST-1", municipio=None):
    # municipio es NOT NULL a nivel de BD (migración 0005_finalizar_
    # municipio_fk), aunque el campo del modelo todavía declare
    # null=True — se arma la cadena completa del catálogo si no
    # se pasa una ya existente.
    if municipio is None:
        _, _, municipio = _crear_cadena_geografica()
    return EstacionServicio.objects.create(
        nombre="Estación de prueba",
        codigo=codigo,
        direccion="Calle Falsa 123",
        municipio=municipio,
    )


class CalcularCupoMensualTests(TestCase):

    def setUp(self):
        self.consumidor = _crear_consumidor("cupo1@test.com")

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_solicitud_aprobada_mes_anterior_no_cuenta_para_mes_siguiente(self, mock_localdate):
        # "Hoy" es 1 de octubre — la solicitud del 30/sept no debe
        # arrastrarse al cupo de octubre (corte de mes calendario).
        mock_localdate.return_value = date(2026, 10, 1)

        fecha_sept = timezone.make_aware(datetime(2026, 9, 30, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA,
            litros_aprobados=100, fecha_aprobacion=fecha_sept,
        )

        cupo = calcular_cupo_mensual(self.consumidor)
        self.assertEqual(cupo["usado"], 0)
        self.assertEqual(cupo["disponible"], 120)
        self.assertEqual(cupo["mes"], "2026-10")

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_rechazada_no_cuenta_al_acumulado(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        fecha = timezone.make_aware(datetime(2026, 9, 10, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.RECHAZADA,
            litros_aprobados=100, fecha_aprobacion=fecha,
        )
        cupo = calcular_cupo_mensual(self.consumidor)
        self.assertEqual(cupo["usado"], 0)

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_vencida_no_cuenta_al_acumulado(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        fecha = timezone.make_aware(datetime(2026, 9, 10, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.EXPIRADA,
            litros_aprobados=100, fecha_aprobacion=fecha,
        )
        cupo = calcular_cupo_mensual(self.consumidor)
        self.assertEqual(cupo["usado"], 0)

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_aprobada_y_despachada_del_mismo_mes_se_suman(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        fecha1 = timezone.make_aware(datetime(2026, 9, 5, 10, 0))
        fecha2 = timezone.make_aware(datetime(2026, 9, 10, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA,
            litros_aprobados=60, fecha_aprobacion=fecha1,
        )
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA,
            litros_aprobados=40, fecha_aprobacion=fecha2,
        )
        cupo = calcular_cupo_mensual(self.consumidor)
        self.assertEqual(cupo["usado"], 100)
        self.assertEqual(cupo["disponible"], 20)


class SolicitudActivaTests(TestCase):

    def setUp(self):
        self.consumidor = _crear_consumidor("activa1@test.com")

    def test_sin_solicitudes_no_hay_activa(self):
        self.assertIsNone(obtener_solicitud_activa(self.consumidor))

    def test_pendiente_es_activa(self):
        s = _crear_solicitud(self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE)
        self.assertEqual(obtener_solicitud_activa(self.consumidor), s)

    def test_despachada_no_es_activa(self):
        _crear_solicitud(self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA)
        self.assertIsNone(obtener_solicitud_activa(self.consumidor))

    def test_constraint_bd_bloquea_segunda_solicitud_activa(self):
        # Garantía dura contra condiciones de carrera — sigue vigente,
        # sin cambios, tal como estaba antes de este fix.
        _crear_solicitud(self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                _crear_solicitud(self.consumidor, Solicitud.EstadoSolicitud.OBSERVADA)

    def test_alerta_repetitividad_en_revision_no_bloquea_creacion(self):
        # EN_REVISION es un valor de ConsumidorPerfil.alerta_repetitividad,
        # no de Solicitud.estado — no debe afectar esta regla. Solo
        # BLOQUEADO bloquea (esta_bloqueado()), comportamiento sin cambios.
        self.consumidor.alerta_repetitividad = ConsumidorPerfil.EstadoAlerta.EN_REVISION
        self.consumidor.save(update_fields=["alerta_repetitividad"])

        self.assertFalse(self.consumidor.esta_bloqueado())
        puede, _ = verificar_puede_crear_solicitud(self.consumidor, 50)
        self.assertTrue(puede)


class VerificarPuedeCrearSolicitudTests(TestCase):

    def setUp(self):
        self.consumidor = _crear_consumidor("crear1@test.com")

    def test_bloquea_si_tiene_solicitud_activa(self):
        _crear_solicitud(self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE)
        puede, mensaje = verificar_puede_crear_solicitud(self.consumidor, 50)
        self.assertFalse(puede)
        self.assertIn("solicitud activa", mensaje)
        self.assertIn("Ya tienes", mensaje)

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_bloquea_si_excede_cupo_sin_historial(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        puede, mensaje = verificar_puede_crear_solicitud(self.consumidor, 150)
        self.assertFalse(puede)
        self.assertIn("cupo mensual", mensaje)
        self.assertIn("Ya tienes", mensaje)

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_bloquea_si_100l_aprobados_y_pide_50(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        fecha = timezone.make_aware(datetime(2026, 9, 5, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA,
            litros_aprobados=100, fecha_aprobacion=fecha,
        )
        puede, mensaje = verificar_puede_crear_solicitud(self.consumidor, 50)
        self.assertFalse(puede)
        self.assertIn("100", mensaje)
        self.assertIn("20", mensaje)

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_permite_dentro_de_cupo_y_sin_activa(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        puede, mensaje = verificar_puede_crear_solicitud(self.consumidor, 100)
        self.assertTrue(puede)
        self.assertEqual(mensaje, "")


class VerificarPuedeAprobarTests(TestCase):

    def setUp(self):
        self.consumidor = _crear_consumidor("aprobar1@test.com")

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_bloquea_si_aprobar_excederia_cupo(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        fecha = timezone.make_aware(datetime(2026, 9, 5, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA,
            litros_aprobados=100, fecha_aprobacion=fecha,
        )
        pendiente = _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE,
            litros_solicitados=50,
        )
        puede, mensaje = verificar_puede_aprobar(pendiente, 50)
        self.assertFalse(puede)
        self.assertIn("cupo mensual", mensaje)

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_permite_si_aprobacion_entra_en_cupo(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        pendiente = _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE,
            litros_solicitados=50,
        )
        puede, mensaje = verificar_puede_aprobar(pendiente, 50)
        self.assertTrue(puede)


class AprobarSolicitudCupoIntegrationTests(TestCase):
    """
    Integración end-to-end del punto 5 de las decisiones aprobadas:
    ANH aprueba una solicitud que excedería el cupo → debe fallar con
    un ValidationError amigable, sin llegar a cambiar el estado.
    """

    def setUp(self):
        self.consumidor = _crear_consumidor(
            "aprobar_full@test.com",
            estado_identidad=ConsumidorPerfil.EstadoIdentidad.VERIFICADO,
        )
        self.anh_user = User.objects.create_user(
            email="anh_full@test.com",
            nombres="ANH",
            apellido_paterno="Test",
            tipo_usuario=User.TipoUsuario.ANH,
            password="testpass123",
        )
        self.estacion = _crear_estacion()

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_anh_aprueba_solicitud_que_excederia_cupo(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        fecha = timezone.make_aware(datetime(2026, 9, 5, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA,
            litros_aprobados=100, fecha_aprobacion=fecha,
        )
        pendiente = _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE,
            litros_solicitados=50,
        )

        with self.assertRaises(DjangoValidationError) as ctx:
            aprobar_solicitud(
                solicitud_id=pendiente.id,
                usuario_aprobador=self.anh_user,
                estacion_servicio=self.estacion,
                litros_aprobados=50,
                tipo_combustible_aprobado=Solicitud.TipoCombustible.GASOLINA,
            )
        self.assertIn("cupo mensual", str(ctx.exception))

        pendiente.refresh_from_db()
        self.assertEqual(pendiente.estado, Solicitud.EstadoSolicitud.PENDIENTE)

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_anh_aprueba_solicitud_dentro_de_cupo(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        pendiente = _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE,
            litros_solicitados=50,
        )

        resultado = aprobar_solicitud(
            solicitud_id=pendiente.id,
            usuario_aprobador=self.anh_user,
            estacion_servicio=self.estacion,
            litros_aprobados=50,
            tipo_combustible_aprobado=Solicitud.TipoCombustible.GASOLINA,
        )
        self.assertEqual(resultado.estado, Solicitud.EstadoSolicitud.APROBADA)


class SolicitudCreateAPITests(TestCase):
    """
    Integración vía API real de SolicitudCreateSerializer.validate() —
    a diferencia de los tests anteriores (que llaman las funciones de
    validar_cupo.py directo), acá se ejercita self.context["request"]
    tal como lo arma DRF en un request HTTP real, para no dar por
    buena esa integración solo por inspección de código.
    """

    def setUp(self):
        self.consumidor = _crear_consumidor("api_create@test.com")
        self.consumidor.user.estado_cuenta = self.consumidor.user.EstadoCuenta.ACTIVO
        self.consumidor.user.save(update_fields=["estado_cuenta"])

        _, _, municipio = _crear_cadena_geografica()
        self.estacion = _crear_estacion(municipio=municipio)
        self.municipio = municipio

        self.client = APIClient()
        self.client.force_authenticate(user=self.consumidor.user)

    def _payload(self, litros=50):
        return {
            "tipo_combustible": Solicitud.TipoCombustible.GASOLINA,
            "litros_solicitados": litros,
            "declaracion_jurada_confirmada": True,
            "departamento": self.municipio.provincia.departamento_id,
            "provincia": self.municipio.provincia_id,
            "municipio": self.municipio.id,
            "estacion_servicio": self.estacion.id,
        }

    def test_bloquea_creacion_si_ya_tiene_activa(self):
        _crear_solicitud(self.consumidor, Solicitud.EstadoSolicitud.PENDIENTE)

        response = self.client.post(
            reverse("solicitud-list"), self._payload(), format="json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("solicitud activa", str(response.data))

    @patch("solicitudes.services.validar_cupo.timezone.localdate")
    def test_bloquea_creacion_si_excede_cupo(self, mock_localdate):
        mock_localdate.return_value = date(2026, 9, 15)
        fecha = timezone.make_aware(datetime(2026, 9, 5, 10, 0))
        _crear_solicitud(
            self.consumidor, Solicitud.EstadoSolicitud.DESPACHADA,
            litros_aprobados=100, fecha_aprobacion=fecha,
        )

        response = self.client.post(
            reverse("solicitud-list"), self._payload(litros=50), format="json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("cupo mensual", str(response.data))

    def test_permite_creacion_dentro_de_cupo_y_sin_activa(self):
        response = self.client.post(
            reverse("solicitud-list"), self._payload(litros=50), format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)


# ------------------------------------------------
# LOTE A — FECHAS EN HORA LOCAL (America/La_Paz, UTC-4)
# ------------------------------------------------

def _utc(*args):
    from datetime import timezone as dt_timezone
    return datetime(*args, tzinfo=dt_timezone.utc)


def _crear_con_fecha(consumidor, fecha_creacion):
    # CANCELADA: no choca con la constraint de "una activa por consumidor".
    # fecha_creacion es auto_now_add, así que se pisa con update().
    s = _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.CANCELADA)
    Solicitud.objects.filter(pk=s.pk).update(fecha_creacion=fecha_creacion)
    return s


class DashboardHoraLocalTests(TestCase):

    def setUp(self):
        self.anh = User.objects.create_user(
            email="anh_dash@test.com", nombres="ANH", apellido_paterno="Dash",
            tipo_usuario=User.TipoUsuario.ANH, password="testpass123",
        )
        self.anh.estado_cuenta = User.EstadoCuenta.ACTIVO
        self.anh.save(update_fields=["estado_cuenta"])
        self.client = APIClient()
        self.client.force_authenticate(user=self.anh)
        self.consumidor = _crear_consumidor("cons_dash@test.com")

    def _dashboard(self, ahora_utc):
        with patch("django.utils.timezone.now", return_value=ahora_utc):
            response = self.client.get(reverse("dashboard-anh"))
        self.assertEqual(response.status_code, 200)
        return response.data

    def test_solicitud_de_las_21_hora_local_cuenta_como_hoy(self):
        # "Ahora": 21:30 del 10/09 en La Paz = 01:30 UTC del 11/09
        _crear_con_fecha(self.consumidor, _utc(2026, 9, 11, 1, 0))   # 21:00 local del 10
        _crear_con_fecha(self.consumidor, _utc(2026, 9, 10, 23, 0))  # 19:00 local del 10
        _crear_con_fecha(self.consumidor, _utc(2026, 9, 10, 1, 0))   # 21:00 local del 09

        data = self._dashboard(_utc(2026, 9, 11, 1, 30))

        # Con la medianoche UTC, "hoy" arrancaba a las 20:00 del 10 y
        # la de las 19:00 quedaba afuera.
        self.assertEqual(data["solicitudes"]["hoy"]["nuevas"], 2)

        tendencia = data["tendencia_7_dias"]
        self.assertEqual(tendencia[-1]["fecha"], "10/09")
        self.assertEqual(tendencia[-1]["solicitudes"], 2)
        self.assertEqual(tendencia[-2]["fecha"], "09/09")
        self.assertEqual(tendencia[-2]["solicitudes"], 1)
        self.assertEqual(data["generado_en"], "10/09/2026 21:30")

    def test_solicitud_de_anoche_no_cuenta_como_hoy_a_la_manana(self):
        _crear_con_fecha(self.consumidor, _utc(2026, 9, 11, 1, 0))   # 21:00 local del 10

        # 10:00 del 11/09 en La Paz = 14:00 UTC del 11/09
        data = self._dashboard(_utc(2026, 9, 11, 14, 0))

        self.assertEqual(data["solicitudes"]["hoy"]["nuevas"], 0)
        self.assertEqual(data["tendencia_7_dias"][-2]["fecha"], "10/09")
        self.assertEqual(data["tendencia_7_dias"][-2]["solicitudes"], 1)


class EstadisticasMesEspanolTests(TestCase):

    def test_evolucion_mensual_con_meses_en_espanol_y_hora_local(self):
        anh = User.objects.create_user(
            email="anh_est@test.com", nombres="ANH", apellido_paterno="Est",
            tipo_usuario=User.TipoUsuario.ANH, password="testpass123",
        )
        anh.estado_cuenta = User.EstadoCuenta.ACTIVO
        anh.save(update_fields=["estado_cuenta"])
        client = APIClient()
        client.force_authenticate(user=anh)

        consumidor = _crear_consumidor("cons_est@test.com")
        # 02:00 UTC del 1° de febrero = 22:00 local del 31 de enero
        _crear_con_fecha(consumidor, _utc(2026, 2, 1, 2, 0))

        # Rango > 31 días → evolución agrupada por mes
        response = client.get("/api/estadisticas/solicitudes/", {
            "fecha_desde": "2026-01-01", "fecha_hasta": "2026-02-28",
        })

        self.assertEqual(response.status_code, 200)
        puntos = response.data["evolucion"]["puntos"]
        self.assertEqual([p["etiqueta"] for p in puntos], ["Ene 2026", "Feb 2026"])
        # Cae en enero (hora local), no en febrero (UTC)
        self.assertEqual([p["creadas"] for p in puntos], [1, 0])


class DeclaracionJuradaFechaTests(TestCase):

    def test_lugar_y_fecha_con_mes_en_espanol(self):
        from reportlab import rl_config
        from solicitudes.services.generar_declaracion_jurada import generar_declaracion_jurada

        consumidor = _crear_consumidor("cons_dj@test.com")
        solicitud  = _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.PENDIENTE)

        # 02:00 UTC del 28/09 = 22:00 local del 27/09. Sin compresión de
        # página el texto queda legible en los bytes del PDF.
        ahora = timezone.localtime(_utc(2026, 9, 28, 2, 0))
        with patch("solicitudes.services.generar_declaracion_jurada.ahora_local", return_value=ahora), \
             patch.object(rl_config, "pageCompression", 0):
            pdf = generar_declaracion_jurada(solicitud)

        self.assertIn(b"27 de septiembre de 2026", pdf)
        self.assertNotIn(b"September", pdf)


class EmailAprobacionHoraLocalTests(TestCase):

    def test_valida_hasta_en_hora_local(self):
        from django.core import mail
        from django.test import override_settings
        from users.email_service import enviar_notificacion_solicitud_aprobada

        consumidor = _crear_consumidor("cons_mail@test.com")
        solicitud  = _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.APROBADA, litros_aprobados=50)
        solicitud.estacion_servicio         = _crear_estacion()
        solicitud.tipo_combustible_aprobado = Solicitud.TipoCombustible.GASOLINA
        solicitud.fecha_expiracion          = _utc(2026, 9, 11, 1, 45)  # 21:45 local del 10

        with override_settings(BREVO_API_KEY=""):
            enviar_notificacion_solicitud_aprobada(solicitud)

        self.assertIn("Válida hasta    : 10/09/2026 21:45", mail.outbox[0].body)


# ------------------------------------------------
# LOTE B — MOTIVO VISIBLE EN EL HISTORIAL DEL CONSUMIDOR (H2)
# ------------------------------------------------

class ListadoExponeMotivoTests(TestCase):

    def test_listado_del_consumidor_incluye_observacion_anh(self):
        consumidor = _crear_consumidor("cons_motivo@test.com")
        consumidor.user.estado_cuenta = User.EstadoCuenta.ACTIVO
        consumidor.user.save(update_fields=["estado_cuenta"])

        s = _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.RECHAZADA)
        Solicitud.objects.filter(pk=s.pk).update(
            observacion_anh="Documento ilegible.",
        )

        client = APIClient()
        client.force_authenticate(user=consumidor.user)
        response = client.get(reverse("solicitud-list"))

        self.assertEqual(response.status_code, 200)
        fila = response.data["results"][0]
        self.assertEqual(fila["estado"], "RECHAZADA")
        self.assertEqual(fila["observacion_anh"], "Documento ilegible.")


class RechazoAutomaticoMotivoTests(TestCase):

    def test_rechazo_por_vencimiento_escribe_motivo_fijo(self):
        from django.test import override_settings
        from solicitudes.services.expirar_solicitudes import rechazar_observadas_vencidas

        consumidor = _crear_consumidor("cons_auto@test.com")
        s = _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.OBSERVADA)
        Solicitud.objects.filter(pk=s.pk).update(
            observacion_anh="Adjunta una foto legible del CI.",
            fecha_limite_respuesta=timezone.now() - timezone.timedelta(hours=1),
        )

        with override_settings(BREVO_API_KEY=""):
            self.assertEqual(rechazar_observadas_vencidas(), 1)

        s.refresh_from_db()
        self.assertEqual(s.estado, Solicitud.EstadoSolicitud.RECHAZADA)
        self.assertEqual(
            s.observacion_anh,
            "Rechazada automáticamente: no se respondió la observación dentro del plazo.",
        )


# ------------------------------------------------
# LOTE D — HISTORIAL POR CONSUMIDOR (ANH/ADMIN)
# ------------------------------------------------

class HistorialPorConsumidorTests(TestCase):

    def test_filtro_por_consumidor_con_estacion(self):
        anh = User.objects.create_user(
            email="anh_hist@test.com", nombres="ANH", apellido_paterno="Hist",
            tipo_usuario=User.TipoUsuario.ANH, password="testpass123",
        )
        anh.estado_cuenta = User.EstadoCuenta.ACTIVO
        anh.save(update_fields=["estado_cuenta"])
        client = APIClient()
        client.force_authenticate(user=anh)

        consumidor = _crear_consumidor("cons_hist@test.com")
        otro       = _crear_consumidor("otro_hist@test.com")
        estacion   = _crear_estacion()

        propia = _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.DESPACHADA)
        Solicitud.objects.filter(pk=propia.pk).update(estacion_servicio=estacion)
        _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.CANCELADA)
        _crear_solicitud(otro, Solicitud.EstadoSolicitud.CANCELADA)

        response = client.get(reverse("solicitud-list"), {"consumidor": consumidor.id})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 2)
        por_id = {r["id_publico"]: r for r in response.data["results"]}
        self.assertEqual(por_id[str(propia.id_publico)]["estacion_nombre"], "Estación de prueba")
        sin_estacion = next(r for r in response.data["results"] if r["id_publico"] != str(propia.id_publico))
        self.assertEqual(sin_estacion["estacion_nombre"], "—")


# ------------------------------------------------
# LOTE E — EMAIL DEL RECHAZO AUTOMÁTICO DESPUÉS DEL COMMIT
# ------------------------------------------------

class RechazoAutomaticoEmailOnCommitTests(TestCase):
    """
    TestCase envuelve cada test en una transacción que nunca se
    confirma, así que los on_commit se inspeccionan con
    captureOnCommitCallbacks.
    """

    def setUp(self):
        consumidor = _crear_consumidor("cons_commit@test.com")
        self.solicitud = _crear_solicitud(consumidor, Solicitud.EstadoSolicitud.OBSERVADA)
        Solicitud.objects.filter(pk=self.solicitud.pk).update(
            observacion_anh="Adjunta una foto legible del CI.",
            fecha_limite_respuesta=timezone.now() - timezone.timedelta(hours=1),
        )

    def test_email_sale_recien_cuando_se_confirma_la_transaccion(self):
        from django.core import mail
        from django.test import override_settings
        from solicitudes.services.expirar_solicitudes import rechazar_observadas_vencidas

        with override_settings(BREVO_API_KEY=""):
            with self.captureOnCommitCallbacks(execute=False) as callbacks:
                self.assertEqual(rechazar_observadas_vencidas(), 1)
                # Dentro de la transacción: todavía no salió nada
                self.assertEqual(len(mail.outbox), 0)

            self.assertEqual(len(callbacks), 1)
            callbacks[0]()  # lo que Django corre al confirmar

        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["cons_commit@test.com"])

    def test_si_la_transaccion_se_revierte_no_sale_email(self):
        from django.core import mail
        from django.test import override_settings
        from solicitudes.services.expirar_solicitudes import rechazar_observadas_vencidas

        class Revertir(Exception):
            pass

        with override_settings(BREVO_API_KEY=""):
            with self.captureOnCommitCallbacks(execute=True) as callbacks:
                with self.assertRaises(Revertir):
                    with transaction.atomic():
                        rechazar_observadas_vencidas()
                        raise Revertir()

        self.assertEqual(callbacks, [])
        self.assertEqual(len(mail.outbox), 0)
        self.solicitud.refresh_from_db()
        self.assertEqual(self.solicitud.estado, Solicitud.EstadoSolicitud.OBSERVADA)


# ------------------------------------------------
# LOTE DEMO — seed_demo
# ------------------------------------------------

class SeedDemoTests(TestCase):

    def setUp(self):
        from io import StringIO
        self.out = StringIO()
        _, _, mun1 = _crear_cadena_geografica()
        self.estacion1 = _crear_estacion("EST-DEMO-A", municipio=mun1)
        self.estacion2 = _crear_estacion("EST-DEMO-B", municipio=mun1)
        self.real = _crear_consumidor("real@ejemplo.com")
        anh = User.objects.create_user(
            email="anh_seed@test.com", nombres="ANH", apellido_paterno="Seed",
            tipo_usuario=User.TipoUsuario.ANH, password="testpass123",
        )
        anh.estado_cuenta = User.EstadoCuenta.ACTIVO
        anh.save(update_fields=["estado_cuenta"])

    def _seed(self, **kwargs):
        from django.core.management import call_command
        call_command("seed_demo", stdout=self.out, **kwargs)

    def _demo(self):
        return Solicitud.objects.filter(consumidor__user__email__endswith="@demo.anh.bo")

    def test_crea_datos_consistentes_sin_emails(self):
        from collections import defaultdict
        from django.core import mail
        from estaciones.models import EstacionServicio

        estaciones_antes = EstacionServicio.objects.count()
        self._seed(password="DemoClave123!")

        usuarios = User.objects.filter(email__endswith="@demo.anh.bo")
        self.assertEqual(usuarios.count(), 15)
        self.assertTrue(all(u.check_password("DemoClave123!") for u in usuarios))
        self.assertEqual(EstacionServicio.objects.count(), estaciones_antes)
        self.assertEqual(len(mail.outbox), 0)

        solicitudes = self._demo()
        self.assertTrue(45 <= solicitudes.count() <= 55, solicitudes.count())
        self.assertEqual(
            set(solicitudes.values_list("estado", flat=True)),
            set(Solicitud.EstadoSolicitud.values),
        )

        ahora = timezone.now()
        for s in solicitudes:
            self.assertLessEqual(s.fecha_creacion, ahora)
            self.assertGreaterEqual(s.fecha_creacion, ahora - timezone.timedelta(days=92))
            if s.estado == "APROBADA":
                self.assertGreater(s.fecha_expiracion, ahora)
            if s.estado == "OBSERVADA":
                self.assertGreater(s.fecha_limite_respuesta, ahora)

        # Una activa por consumidor y ≤ 120 L por mes (APROBADA + DESPACHADA)
        activas  = defaultdict(int)
        por_mes  = defaultdict(int)
        for s in solicitudes:
            if s.estado in ("PENDIENTE", "OBSERVADA", "APROBADA"):
                activas[s.consumidor_id] += 1
            if s.estado in ("APROBADA", "DESPACHADA"):
                local = timezone.localtime(s.fecha_aprobacion)
                por_mes[(s.consumidor_id, local.year, local.month)] += s.litros_aprobados
        self.assertTrue(all(n == 1 for n in activas.values()))
        self.assertTrue(all(litros <= 120 for litros in por_mes.values()), dict(por_mes))

    def test_no_duplica_si_ya_hay_demo(self):
        from django.core.management.base import CommandError
        self._seed(password="DemoClave123!")
        with self.assertRaises(CommandError):
            self._seed(password="DemoClave123!")

    def test_exige_password(self):
        from django.core.management.base import CommandError
        with self.assertRaises(CommandError):
            self._seed()

    def test_limpiar_borra_solo_la_demo(self):
        from solicitudes.models import AuditoriaEstadoSolicitud
        self._seed(password="DemoClave123!")
        real_solicitud = _crear_solicitud(self.real, Solicitud.EstadoSolicitud.PENDIENTE)

        self._seed(limpiar=True)

        self.assertFalse(User.objects.filter(email__endswith="@demo.anh.bo").exists())
        self.assertFalse(self._demo().exists())
        self.assertFalse(AuditoriaEstadoSolicitud.objects.filter(
            solicitud__consumidor__user__email__endswith="@demo.anh.bo").exists())
        self.assertTrue(User.objects.filter(email="real@ejemplo.com").exists())
        self.assertTrue(Solicitud.objects.filter(pk=real_solicitud.pk).exists())


class EmailDemoOmitidoTests(TestCase):

    def test_no_se_envia_a_direcciones_demo(self):
        from django.core import mail
        from django.test import override_settings
        from users.email_service import _enviar_email

        with override_settings(BREVO_API_KEY=""):
            self.assertTrue(_enviar_email("Asunto", "Cuerpo", "juan.mamani.01@demo.anh.bo"))
            self.assertEqual(len(mail.outbox), 0)
            _enviar_email("Asunto", "Cuerpo", "real@ejemplo.com")
            self.assertEqual(len(mail.outbox), 1)


class SeedDemoCITests(TestCase):

    def test_ci_numericos_de_7_a_8_digitos_sin_chocar_con_existentes(self):
        import random
        from io import StringIO
        from django.core.management import call_command
        from consumidores.models import DocumentoIdentidad

        _, _, mun = _crear_cadena_geografica()
        _crear_estacion("EST-CI", municipio=mun)

        # Ocupa de antemano el primer CI que sortearía el seed (semilla
        # por defecto 2026 → generador de CI con 2027).
        primer_ci = str(random.Random(2027).randint(1_000_000, 99_999_999))
        real = _crear_consumidor("real_ci@ejemplo.com")
        DocumentoIdentidad.objects.create(
            perfil=real, tipo_documento="CI", numero_documento=primer_ci,
            anverso="", reverso="",
        )

        call_command("seed_demo", password="DemoClave123!", stdout=StringIO())

        demo = DocumentoIdentidad.objects.filter(perfil__user__email__endswith="@demo.anh.bo")
        numeros = list(demo.values_list("numero_documento", flat=True))
        self.assertEqual(len(numeros), 15)
        self.assertEqual(len(set(numeros)), 15)
        self.assertTrue(all(n.isdigit() and 7 <= len(n) <= 8 for n in numeros), numeros)
        self.assertNotIn(primer_ci, numeros)
        self.assertEqual(
            DocumentoIdentidad.objects.get(numero_documento=primer_ci).perfil, real,
        )


# ------------------------------------------------
# LOTE C1 — REPORTES: ESTADÍSTICAS Y REPORTE DE CONSUMIDORES
# ------------------------------------------------

def _anh_client(email):
    anh = User.objects.create_user(
        email=email, nombres="ANH", apellido_paterno="Rep",
        tipo_usuario=User.TipoUsuario.ANH, password="testpass123",
    )
    anh.estado_cuenta = User.EstadoCuenta.ACTIVO
    anh.save(update_fields=["estado_cuenta"])
    client = APIClient()
    client.force_authenticate(user=anh)
    return client


def _solicitud_en(consumidor, estado, creada, **campos):
    """Solicitud con fecha_creacion fija (hora local) y campos extra."""
    s = _crear_solicitud(consumidor, estado)
    Solicitud.objects.filter(pk=s.pk).update(
        fecha_creacion=timezone.make_aware(creada), **campos,
    )
    return s


class EstadisticasReportesTests(TestCase):

    URL = "/api/estadisticas/solicitudes/"

    def setUp(self):
        self.client = _anh_client("anh_est_c1@test.com")
        c1 = _crear_consumidor("c1_est@test.com")
        c2 = _crear_consumidor("c2_est@test.com")
        aprob = timezone.make_aware(datetime(2026, 8, 5, 15, 0))
        _solicitud_en(c1, "DESPACHADA", datetime(2026, 8, 5, 10, 0),
                      fecha_aprobacion=aprob, litros_aprobados=40, litros_despachados=40)
        _solicitud_en(c1, "EXPIRADA", datetime(2026, 8, 5, 11, 0),
                      fecha_aprobacion=aprob, litros_aprobados=30)
        _solicitud_en(c2, "RECHAZADA", datetime(2026, 8, 10, 9, 0))
        _solicitud_en(c2, "CANCELADA", datetime(2026, 8, 10, 12, 0))

    def _get(self, **params):
        return self.client.get(self.URL, params)

    def test_aprobadas_alguna_vez_incluye_despachadas_y_expiradas_y_tasas(self):
        r = self._get(fecha_desde="2026-08-01", fecha_hasta="2026-08-31")
        self.assertEqual(r.status_code, 200)
        d = r.data
        self.assertEqual(d["total"], 4)
        self.assertEqual(d["aprobadas_alguna_vez"], 2)
        self.assertEqual(d["rechazadas"], 1)
        self.assertEqual(d["resueltas"], 3)
        self.assertEqual(d["tasa_aprobacion"], 66.7)
        self.assertEqual(d["tasa_rechazo"], 33.3)
        self.assertEqual(d["litros_despachados"], 40)
        self.assertEqual(d["periodo"]["texto"], "Solicitudes creadas entre 01/08/2026 y 31/08/2026")
        self.assertEqual(
            {c["tipo"]: c["litros"] for c in d["despachados_por_combustible"]},
            {"GASOLINA": 40, "DIESEL": 0},
        )

    def test_tasas_sin_resueltas_son_none(self):
        r = self._get(fecha_desde="2026-08-01", fecha_hasta="2026-08-31", estado="CANCELADA")
        self.assertEqual(r.data["resueltas"], 0)
        self.assertIsNone(r.data["tasa_aprobacion"])
        self.assertIsNone(r.data["tasa_rechazo"])

    def test_rango_invalido_400(self):
        r = self._get(fecha_desde="2026-08-31", fecha_hasta="2026-08-01")
        self.assertEqual(r.status_code, 400)
        self.assertIn("posterior", str(r.data))

    def test_rango_sin_datos_200_con_total_cero(self):
        r = self._get(fecha_desde="2025-01-01", fecha_hasta="2025-01-31")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["total"], 0)

    def test_agrupacion_por_dia_hasta_31_dias(self):
        evol = self._get(fecha_desde="2026-08-01", fecha_hasta="2026-08-31").data["evolucion"]
        self.assertEqual(evol["agrupacion"], "dia")
        self.assertEqual(len(evol["puntos"]), 31)
        dia5 = next(p for p in evol["puntos"] if p["etiqueta"] == "05/08")
        self.assertEqual((dia5["creadas"], dia5["aprobadas"], dia5["despachadas"]), (2, 2, 1))

    def test_agrupacion_por_mes_si_el_rango_es_mayor(self):
        evol = self._get(fecha_desde="2026-06-01", fecha_hasta="2026-08-31").data["evolucion"]
        self.assertEqual(evol["agrupacion"], "mes")
        self.assertEqual([p["etiqueta"] for p in evol["puntos"]], ["Jun 2026", "Jul 2026", "Ago 2026"])
        self.assertEqual(evol["puntos"][-1]["creadas"], 4)
        self.assertEqual(evol["puntos"][-1]["aprobadas"], 2)

    def test_reporte_de_solicitudes_rechaza_rango_invalido_y_vacio(self):
        url = "/api/reportes/solicitudes/"
        self.assertEqual(self.client.get(url, {"fecha_desde": "2026-08-31", "fecha_hasta": "2026-08-01"}).status_code, 400)
        vacio = self.client.get(url, {"fecha_desde": "2025-01-01", "fecha_hasta": "2025-01-31"})
        self.assertEqual(vacio.status_code, 400)
        self.assertIn("No hay solicitudes", str(vacio.data))
        ok = self.client.get(url, {"fecha_desde": "2026-08-01", "fecha_hasta": "2026-08-31"})
        self.assertEqual(ok.status_code, 200)


class ReporteConsumidoresMesTests(TestCase):

    def setUp(self):
        from solicitudes.services.validar_cupo import calcular_cupo_mensual
        self.calcular = calcular_cupo_mensual
        self.agotado = _crear_consumidor("agotado@test.com")
        self.parcial = _crear_consumidor("parcial@test.com")
        aprob = timezone.make_aware(datetime(2026, 8, 10, 12, 0))
        desp  = timezone.make_aware(datetime(2026, 8, 11, 12, 0))
        _solicitud_en(self.agotado, "DESPACHADA", datetime(2026, 8, 10, 8, 0),
                      fecha_aprobacion=aprob, fecha_despacho=desp,
                      litros_aprobados=120, litros_despachados=120)
        _solicitud_en(self.parcial, "DESPACHADA", datetime(2026, 8, 12, 8, 0),
                      fecha_aprobacion=aprob, fecha_despacho=desp,
                      litros_aprobados=50, litros_despachados=50)

    def test_cupo_agotado_por_mes_con_la_logica_de_validar_cupo(self):
        from solicitudes.services.generar_reportes import get_consumidores

        agosto = get_consumidores("CUPO_AGOTADO", date(2026, 8, 1))
        self.assertEqual([p.pk for p in agosto], [self.agotado.pk])
        self.assertEqual(get_consumidores("CUPO_AGOTADO", date(2026, 9, 1)), [])

        # Mismo número que calcular_cupo_mensual para ese mes
        for perfil in get_consumidores("TODOS", date(2026, 8, 1)):
            self.assertEqual(perfil.cupo_usado_mes, self.calcular(perfil, mes=date(2026, 8, 20))["usado"])

    def test_ordenado_por_litros_despachados(self):
        from solicitudes.services.generar_reportes import fila_consumidor, get_consumidores
        filas = [fila_consumidor(p) for p in get_consumidores("TODOS", date(2026, 8, 1))]
        self.assertEqual([f["litros_despachados"] for f in filas], [120, 50])
        self.assertEqual(filas[0]["cupo_texto"], "120/120 L")
        self.assertEqual(filas[0]["solicitudes_mes"], 1)

    def test_sin_mes_calcular_cupo_mantiene_el_mes_actual(self):
        # La solicitud de agosto no cuenta para el mes actual (septiembre)
        with patch("solicitudes.services.validar_cupo.timezone.localdate", return_value=date(2026, 9, 15)):
            self.assertEqual(self.calcular(self.agotado)["usado"], 0)

    def test_endpoint_400_sin_datos_y_filtro_invalido(self):
        client = _anh_client("anh_cons_c1@test.com")
        url = "/api/reportes/consumidores/"
        self.assertEqual(client.get(url, {"filtro": "CUPO_AGOTADO", "mes": "2026-09"}).status_code, 400)
        self.assertEqual(client.get(url, {"filtro": "SUPERARON_LIMITE", "mes": "2026-08"}).status_code, 400)
        # Sin consumidores bloqueados → 400 (no se generan archivos vacíos)
        self.assertEqual(client.get(url, {"filtro": "BLOQUEADOS", "mes": "2026-08"}).status_code, 400)
        self.assertEqual(client.get(url, {"filtro": "CUPO_AGOTADO", "mes": "2026-08"}).status_code, 200)

    def test_reporte_sin_n_mas_1(self):
        from solicitudes.services.generar_reportes import generar_reporte_excel, generar_reporte_pdf
        mes = date(2026, 8, 1)
        with self.assertNumQueries(2):
            generar_reporte_excel("TODOS", mes)
        with self.assertNumQueries(2):
            generar_reporte_pdf("TODOS", mes)

        # Más consumidores, mismas queries
        for i in range(5):
            c = _crear_consumidor(f"extra{i}@test.com")
            _solicitud_en(c, "CANCELADA", datetime(2026, 8, 3, 8, 0))
        with self.assertNumQueries(2):
            generar_reporte_excel("TODOS", mes)
        with self.assertNumQueries(2):
            generar_reporte_pdf("TODOS", mes)


# ------------------------------------------------
# LOTE C2 — FORMATO DE ARCHIVOS DE REPORTES
# ------------------------------------------------

def _pdf_legible(generar, *args, **kwargs):
    """PDF sin compresión de página, para buscar texto en los bytes."""
    from reportlab import rl_config
    with patch.object(rl_config, "pageCompression", 0):
        return generar(*args, **kwargs)


def _libro(contenido):
    import io
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(contenido))


class ReportesArchivosC2Tests(TestCase):

    AGOSTO = date(2026, 8, 1)

    def setUp(self):
        _, _, mun = _crear_cadena_geografica()
        self.estacion = _crear_estacion("EST-C2", municipio=mun)
        aprob = timezone.make_aware(datetime(2026, 8, 10, 12, 0))
        desp  = timezone.make_aware(datetime(2026, 8, 11, 12, 0))

        self.a = _crear_consumidor("a_c2@test.com")
        self.b = _crear_consumidor("b_c2@test.com")
        _solicitud_en(self.a, "DESPACHADA", datetime(2026, 8, 10, 8, 0),
                      estacion_servicio=self.estacion, tipo_combustible_aprobado="GASOLINA",
                      fecha_aprobacion=aprob, fecha_despacho=desp,
                      litros_aprobados=120, litros_despachados=120)
        _solicitud_en(self.b, "DESPACHADA", datetime(2026, 8, 12, 8, 0),
                      estacion_servicio=self.estacion, tipo_combustible="DIESEL",
                      tipo_combustible_aprobado="DIESEL",
                      fecha_aprobacion=aprob, fecha_despacho=desp,
                      litros_aprobados=50, litros_despachados=50)
        _solicitud_en(self.b, "RECHAZADA", datetime(2026, 8, 20, 8, 0))

        self.bloqueado = _crear_consumidor("bloq_c2@test.com")
        ConsumidorPerfil.objects.filter(pk=self.bloqueado.pk).update(
            alerta_repetitividad="BLOQUEADO", motivo_bloqueo="Uso indebido del cupo.",
        )
        self.revision = _crear_consumidor("rev_c2@test.com")
        ConsumidorPerfil.objects.filter(pk=self.revision.pk).update(alerta_repetitividad="EN_REVISION")

        self.filtros = {"desde": date(2026, 8, 1), "hasta": date(2026, 8, 31),
                        "estado": None, "combustible": None, "estacion": None}

    # ---------------- consumidores

    def test_filtros_por_alerta_con_columna_motivo(self):
        from solicitudes.services.generar_reportes import cabeceras, fila_consumidor, get_consumidores

        bloq = get_consumidores("BLOQUEADOS", self.AGOSTO)
        self.assertEqual([p.pk for p in bloq], [self.bloqueado.pk])
        self.assertEqual(fila_consumidor(bloq[0])["motivo"], "Uso indebido del cupo.")
        self.assertEqual([p.pk for p in get_consumidores("EN_REVISION", self.AGOSTO)], [self.revision.pk])
        self.assertIn("Motivo", cabeceras("BLOQUEADOS"))
        self.assertNotIn("Motivo", cabeceras("TODOS"))

        wb = _libro(__import__("solicitudes.services.generar_reportes", fromlist=["x"])
                    .generar_reporte_excel("BLOQUEADOS", self.AGOSTO))
        ws = wb["Consumidores"]
        self.assertEqual(ws.cell(row=1, column=8).value, "Motivo")
        self.assertEqual(ws.cell(row=2, column=8).value, "Uso indebido del cupo.")

    def test_excel_consumidores_con_y_sin_detalle(self):
        from solicitudes.services.generar_reportes import generar_reporte_excel

        sin = _libro(generar_reporte_excel("TODOS", self.AGOSTO))
        self.assertEqual(sin.sheetnames, ["Resumen", "Consumidores"])

        con = _libro(generar_reporte_excel("TODOS", self.AGOSTO, incluir_detalle=True))
        self.assertEqual(con.sheetnames, ["Resumen", "Consumidores", "Detalle"])
        det = con["Detalle"]
        self.assertEqual(det.max_row - 1, 3)                       # 3 solicitudes de agosto
        self.assertEqual(det.freeze_panes, "A2")
        self.assertIsNotNone(det.auto_filter.ref)
        self.assertIsInstance(det.cell(row=2, column=4).value, datetime)   # fecha real
        self.assertIsInstance(det.cell(row=2, column=6).value, int)        # litros número

    def test_pdf_consumidores_detalle_opcional_y_nota(self):
        from solicitudes.services.generar_reportes import generar_reporte_pdf

        sin = _pdf_legible(generar_reporte_pdf, "TODOS", self.AGOSTO)
        con = _pdf_legible(generar_reporte_pdf, "TODOS", self.AGOSTO, incluir_detalle=True)
        self.assertNotIn(b"Detalle de solicitudes por consumidor", sin)
        self.assertIn(b"Detalle de solicitudes por consumidor", con)
        self.assertIn(b"regla de 120 L", sin)
        self.assertIn(b"gina 1 de", sin)   # "Página 1 de N" (la á va codificada)

    def test_queries_fijas_con_detalle(self):
        from solicitudes.services.generar_reportes import generar_reporte_excel, generar_reporte_pdf

        with self.assertNumQueries(3):
            generar_reporte_excel("TODOS", self.AGOSTO, incluir_detalle=True)
        with self.assertNumQueries(3):
            generar_reporte_pdf("TODOS", self.AGOSTO, incluir_detalle=True)

        for i in range(4):
            c = _crear_consumidor(f"extra_c2_{i}@test.com")
            _solicitud_en(c, "CANCELADA", datetime(2026, 8, 5, 8, 0))
        with self.assertNumQueries(3):
            generar_reporte_excel("TODOS", self.AGOSTO, incluir_detalle=True)
        with self.assertNumQueries(3):
            generar_reporte_pdf("TODOS", self.AGOSTO, incluir_detalle=True)

    # ---------------- solicitudes

    def test_excel_solicitudes_hojas_y_resumen_por_estacion(self):
        from solicitudes.services.generar_reportes_solicitudes import generar_excel_solicitudes

        wb = _libro(generar_excel_solicitudes(self.filtros))
        self.assertEqual(wb.sheetnames, ["Resumen", "Por estación", "Detalle"])

        est = wb["Por estación"]
        self.assertEqual([c.value for c in est[2]], ["Estación de prueba", 2, 120, 50, 170])

        det = wb["Detalle"]
        self.assertEqual(det.freeze_panes, "A2")
        self.assertEqual(det.max_row - 1, 3)
        self.assertIsInstance(det.cell(row=2, column=10).value, datetime)
        self.assertIsInstance(det.cell(row=2, column=4).value, int)

    def test_pdf_solicitudes_con_filtros_legibles_y_resumen_por_estacion(self):
        from solicitudes.services.generar_reportes_solicitudes import generar_pdf_solicitudes

        filtros = {**self.filtros, "estado": "DESPACHADA", "estacion": self.estacion.id}
        pdf = _pdf_legible(generar_pdf_solicitudes, filtros)
        self.assertIn(b"Estado: Despachada", pdf)
        # reportlab escribe la "ó" como escape octal (\363) en el PDF
        self.assertIn(b"Estaci\\363n: Estaci\\363n de prueba", pdf)
        self.assertIn(b"Resumen por estaci", pdf)
        self.assertNotIn(b"Estaci\\363n ID", pdf)

    def test_archivos_con_un_solo_registro(self):
        from solicitudes.services.generar_reportes import generar_reporte_excel, generar_reporte_pdf
        from solicitudes.services.generar_reportes_solicitudes import (
            generar_excel_solicitudes, generar_pdf_solicitudes,
        )
        uno = {**self.filtros, "estado": "RECHAZADA"}
        self.assertTrue(generar_pdf_solicitudes(uno).startswith(b"%PDF"))
        self.assertEqual(_libro(generar_excel_solicitudes(uno))["Detalle"].max_row, 2)
        self.assertTrue(generar_reporte_pdf("BLOQUEADOS", self.AGOSTO, incluir_detalle=True).startswith(b"%PDF"))
        self.assertEqual(_libro(generar_reporte_excel("BLOQUEADOS", self.AGOSTO))["Consumidores"].max_row, 2)

    # ---------------- nombres de archivo

    def test_nombre_de_archivo_con_el_periodo(self):
        client = _anh_client("anh_c2@test.com")

        r = client.get("/api/reportes/solicitudes/", {
            "fecha_desde": "2026-08-01", "fecha_hasta": "2026-08-31", "formato": "PDF",
        })
        self.assertEqual(r.status_code, 200)
        self.assertIn('filename="reporte_solicitudes_2026-08-01_2026-08-31.pdf"', r["Content-Disposition"])

        r = client.get("/api/reportes/consumidores/", {"mes": "2026-08", "formato": "EXCEL"})
        self.assertIn('filename="reporte_consumidores_2026-08.xlsx"', r["Content-Disposition"])

        r = client.get("/api/reportes/consumidores/", {
            "mes": "2026-08", "formato": "EXCEL", "filtro": "CUPO_AGOTADO", "incluir_detalle": "true",
        })
        self.assertEqual(r.status_code, 200)
        self.assertIn('filename="reporte_consumidores_2026-08_cupo_agotado.xlsx"', r["Content-Disposition"])
        self.assertIn("Detalle", _libro(r.content).sheetnames)


class LitrosDespachadosCoincidenTests(TestCase):
    """
    Consumidores y solicitudes usan la misma base para los litros
    despachados: solicitudes CREADAS en el mes y despachadas.
    """

    def test_total_de_litros_despachados_coincide_entre_reportes(self):
        from solicitudes.services.estadisticas import calcular_indicadores, solicitudes_filtradas
        from solicitudes.services.generar_reportes import get_consumidores

        a = _crear_consumidor("coinc_a@test.com")
        b = _crear_consumidor("coinc_b@test.com")
        aware = timezone.make_aware

        # Creada y despachada en agosto → cuenta
        _solicitud_en(a, "DESPACHADA", datetime(2026, 8, 5, 9, 0),
                      fecha_aprobacion=aware(datetime(2026, 8, 5, 12, 0)),
                      fecha_despacho=aware(datetime(2026, 8, 6, 9, 0)),
                      litros_aprobados=40, litros_despachados=40)
        # Creada en agosto, despachada en septiembre → cuenta para agosto
        _solicitud_en(b, "DESPACHADA", datetime(2026, 8, 30, 9, 0),
                      fecha_aprobacion=aware(datetime(2026, 8, 30, 12, 0)),
                      fecha_despacho=aware(datetime(2026, 9, 1, 9, 0)),
                      litros_aprobados=25, litros_despachados=25)
        # Creada en julio, despachada en agosto → NO cuenta para agosto
        _solicitud_en(b, "DESPACHADA", datetime(2026, 7, 30, 9, 0),
                      fecha_aprobacion=aware(datetime(2026, 7, 30, 12, 0)),
                      fecha_despacho=aware(datetime(2026, 8, 2, 9, 0)),
                      litros_aprobados=30, litros_despachados=30)

        agosto = {"desde": date(2026, 8, 1), "hasta": date(2026, 8, 31),
                  "estado": None, "combustible": None, "estacion": None}
        total_solicitudes = calcular_indicadores(solicitudes_filtradas(agosto))["litros_despachados"]
        total_consumidores = sum(
            p.litros_despachados_mes for p in get_consumidores("TODOS", date(2026, 8, 1))
        )

        self.assertEqual(total_solicitudes, 65)
        self.assertEqual(total_consumidores, total_solicitudes)


# ------------------------------------------------
# LOTE C3 — REPORTE POR CONSUMIDOR
# ------------------------------------------------

class ReportePorConsumidorTests(TestCase):

    URL = "/api/reportes/consumidores/"

    def setUp(self):
        from consumidores.models import DocumentoIdentidad
        self.client = _anh_client("anh_c3@test.com")
        self.uno  = _crear_consumidor("uno_c3@test.com")
        self.otro = _crear_consumidor("otro_c3@test.com")
        DocumentoIdentidad.objects.create(
            perfil=self.uno, tipo_documento="CI", numero_documento="7654321",
            anverso="", reverso="",
        )
        aprob = timezone.make_aware(datetime(2026, 8, 10, 12, 0))
        desp  = timezone.make_aware(datetime(2026, 8, 11, 12, 0))
        for perfil in (self.uno, self.otro):
            _solicitud_en(perfil, "DESPACHADA", datetime(2026, 8, 10, 8, 0),
                          fecha_aprobacion=aprob, fecha_despacho=desp,
                          litros_aprobados=40, litros_despachados=40)
        self.sin_datos = _crear_consumidor("vacio_c3@test.com")

    def test_filtra_por_consumidor_y_fuerza_detalle(self):
        from solicitudes.services.generar_reportes import get_consumidores

        self.assertEqual(
            [p.pk for p in get_consumidores("TODOS", date(2026, 8, 1), consumidor_id=self.uno.pk)],
            [self.uno.pk],
        )

        r = self.client.get(self.URL, {"mes": "2026-08", "formato": "EXCEL",
                                       "consumidor_id": self.uno.pk, "filtro": "CUPO_AGOTADO"})
        self.assertEqual(r.status_code, 200)
        self.assertIn('filename="reporte_consumidor_7654321_2026-08.xlsx"', r["Content-Disposition"])
        wb = _libro(r.content)
        # Detalle forzado aunque no se pidió; filtro ignorado (no está agotado)
        self.assertEqual(wb.sheetnames, ["Resumen", "Consumidores", "Detalle"])
        self.assertEqual(wb["Consumidores"].max_row, 2)
        self.assertEqual(wb["Detalle"].max_row, 2)

        pdf = self.client.get(self.URL, {"mes": "2026-08", "formato": "PDF", "consumidor_id": self.uno.pk})
        self.assertEqual(pdf.status_code, 200)
        self.assertIn('filename="reporte_consumidor_7654321_2026-08.pdf"', pdf["Content-Disposition"])

    def test_consumidor_inexistente_404(self):
        r = self.client.get(self.URL, {"mes": "2026-08", "consumidor_id": 999999})
        self.assertEqual(r.status_code, 404)

    def test_consumidor_sin_solicitudes_en_el_mes_400(self):
        r = self.client.get(self.URL, {"mes": "2026-08", "consumidor_id": self.sin_datos.pk})
        self.assertEqual(r.status_code, 400)
        r = self.client.get(self.URL, {"mes": "2026-07", "consumidor_id": self.uno.pk})
        self.assertEqual(r.status_code, 400)

    def test_queries_fijas_por_consumidor(self):
        from solicitudes.services.generar_reportes import generar_reporte_excel, generar_reporte_pdf
        mes = date(2026, 8, 1)
        with self.assertNumQueries(3):
            generar_reporte_pdf("TODOS", mes, consumidor_id=self.uno.pk)
        with self.assertNumQueries(3):
            generar_reporte_excel("TODOS", mes, consumidor_id=self.uno.pk)
