# apps/solicitudes/services/generar_reportes.py
#
# Reporte de consumidores de un MES calendario (hora local).
#
#   - TODOS        → todos los consumidores.
#   - CUPO_AGOTADO → cupo usado >= 120 L en ese mes, con la misma regla
#                    que validar_cupo (APROBADA + DESPACHADA por mes de
#                    fecha_aprobacion).
#
# Por consumidor: solicitudes creadas en el mes (cualquier estado),
# litros despachados en el mes (por fecha de despacho) y cupo usado.
# Todo sale de un único queryset anotado + prefetch: la cantidad de
# queries no crece con la cantidad de consumidores.

import io
from datetime import date

from django.db.models import Count, IntegerField, Prefetch, Q, Sum, Value
from django.db.models.functions import Coalesce

from core.fechas import (
    ahora_local, formatear_fecha, formatear_fecha_hora,
    formatear_mes_anio_largo, rango_mes_local,
)


FILTROS = {
    "TODOS":        "Todos los consumidores",
    "CUPO_AGOTADO": "Cupo mensual agotado",
}


# ------------------------------------------------
# DATOS
# ------------------------------------------------

def get_consumidores(filtro: str, mes: date, con_solicitudes: bool = False):
    """
    Lista de ConsumidorPerfil del reporte, ordenada por litros
    despachados (mayor a menor). Cada perfil trae anotados
    `solicitudes_mes`, `litros_despachados_mes` y `cupo_usado_mes`, y
    con `con_solicitudes` también `solicitudes_del_mes` (lista).
    """
    from consumidores.models import ConsumidorPerfil
    from solicitudes.models import Solicitud
    from solicitudes.services.validar_cupo import CUPO_MENSUAL_LITROS, filtro_cupo_del_mes

    inicio, fin = rango_mes_local(mes)
    cero = Value(0, output_field=IntegerField())

    qs = (
        ConsumidorPerfil.objects
        .select_related("user", "municipio")
        .prefetch_related("documentos")
        .annotate(
            solicitudes_mes=Count(
                "solicitudes",
                filter=Q(solicitudes__fecha_creacion__gte=inicio,
                         solicitudes__fecha_creacion__lt=fin),
            ),
            litros_despachados_mes=Coalesce(
                Sum("solicitudes__litros_despachados",
                    filter=Q(solicitudes__estado=Solicitud.EstadoSolicitud.DESPACHADA,
                             solicitudes__fecha_despacho__gte=inicio,
                             solicitudes__fecha_despacho__lt=fin)),
                cero,
            ),
            cupo_usado_mes=Coalesce(
                Sum("solicitudes__litros_aprobados",
                    filter=filtro_cupo_del_mes(inicio, fin, prefijo="solicitudes__")),
                cero,
            ),
        )
        .order_by("-litros_despachados_mes", "user__apellido_paterno", "user__nombres")
    )

    if filtro == "CUPO_AGOTADO":
        qs = qs.filter(cupo_usado_mes__gte=CUPO_MENSUAL_LITROS)

    if con_solicitudes:
        qs = qs.prefetch_related(Prefetch(
            "solicitudes",
            queryset=Solicitud.objects
                .filter(fecha_creacion__gte=inicio, fecha_creacion__lt=fin)
                .select_related("estacion_servicio")
                .order_by("fecha_creacion"),
            to_attr="solicitudes_del_mes",
        ))

    return list(qs)


def _ci(perfil) -> str:
    # documentos ya viene prefetcheado: .all() usa la caché (.first()
    # dispararía una query por consumidor).
    docs = perfil.documentos.all()
    if not docs:
        return "—"
    doc = docs[0]
    comp = f"-{doc.complemento_documento}" if doc.complemento_documento else ""
    return f"{doc.tipo_documento} {doc.numero_documento}{comp}"


def fila_consumidor(perfil) -> dict:
    from solicitudes.services.validar_cupo import CUPO_MENSUAL_LITROS
    return {
        "nombre":             perfil.user.nombre_completo(),
        "ci":                 _ci(perfil),
        "municipio":          perfil.municipio.nombre if perfil.municipio else "—",
        "solicitudes_mes":    perfil.solicitudes_mes,
        "litros_despachados": perfil.litros_despachados_mes,
        "cupo_usado":         perfil.cupo_usado_mes,
        "cupo_texto":         f"{perfil.cupo_usado_mes}/{CUPO_MENSUAL_LITROS} L",
        "estado_cuenta":      perfil.user.get_estado_cuenta_display(),
    }


def titulo_reporte(filtro: str, mes: date) -> str:
    return f"Reporte de consumidores — {FILTROS[filtro]} — {formatear_mes_anio_largo(mes)}"


CABECERAS = [
    "Nombre", "CI", "Municipio", "Solicitudes del mes",
    "Litros despachados", "Cupo usado", "Estado de cuenta",
]


def _valores(d: dict) -> list:
    return [
        d["nombre"], d["ci"], d["municipio"], d["solicitudes_mes"],
        d["litros_despachados"], d["cupo_texto"], d["estado_cuenta"],
    ]


# ------------------------------------------------
# EXCEL
# ------------------------------------------------

def generar_reporte_excel(filtro: str, mes: date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Consumidores"

    header_fill    = PatternFill("solid", fgColor="1a3a5c")
    subheader_fill = PatternFill("solid", fgColor="2d6a9f")
    thin = Side(style="thin")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(f"A1:{get_column_letter(len(CABECERAS))}1")
    ws["A1"] = f"{titulo_reporte(filtro, mes)} — Generado: {formatear_fecha_hora(ahora_local())}"
    ws["A1"].font      = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill      = header_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    for col, cab in enumerate(CABECERAS, 1):
        c = ws.cell(row=2, column=col, value=cab)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = subheader_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde

    for fila, perfil in enumerate(get_consumidores(filtro, mes), 3):
        for col, valor in enumerate(_valores(fila_consumidor(perfil)), 1):
            c = ws.cell(row=fila, column=col, value=valor)
            c.border = borde

    for i, ancho in enumerate([30, 18, 18, 12, 14, 12, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------
# PDF
# ------------------------------------------------

def generar_reporte_pdf(filtro: str, mes: date) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        rightMargin=1.5 * cm, leftMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )

    azul, azul2 = colors.HexColor("#1a3a5c"), colors.HexColor("#2d6a9f")
    gris, blanco = colors.HexColor("#f2f2f2"), colors.white

    titulo_style = ParagraphStyle("titulo", fontSize=14, textColor=blanco,
                                  alignment=TA_CENTER, fontName="Helvetica-Bold")
    subtitulo_style = ParagraphStyle("subtitulo", fontSize=11, textColor=azul,
                                     alignment=TA_CENTER, fontName="Helvetica-Bold", spaceAfter=6)
    bold_style = ParagraphStyle("bold", fontSize=9, fontName="Helvetica-Bold", textColor=azul)

    elementos = []

    encabezado = Table([[Paragraph(
        "AGENCIA NACIONAL DE HIDROCARBUROS — BOLIVIA<br/>"
        f"{titulo_reporte(filtro, mes)}<br/>"
        f"<font size=9>Generado: {formatear_fecha_hora(ahora_local())}</font>",
        titulo_style,
    )]], colWidths=["100%"])
    encabezado.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), azul),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    elementos += [encabezado, Spacer(1, 0.4 * cm)]

    consumidores = get_consumidores(filtro, mes, con_solicitudes=True)

    elementos.append(Paragraph("RESUMEN DE CONSUMIDORES", subtitulo_style))
    filas = [CABECERAS] + [
        [str(v) for v in _valores(fila_consumidor(p))] for p in consumidores
    ]
    tabla = Table(filas, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), azul2),
        ("TEXTCOLOR", (0, 0), (-1, 0), blanco),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [blanco, gris]),
        ("ALIGN", (3, 1), (5, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BOX", (0, 0), (-1, -1), 1, azul),
    ]))
    elementos += [tabla, Spacer(1, 0.5 * cm)]

    # Detalle: solicitudes del mes de cada consumidor (desde el prefetch)
    con_detalle = [p for p in consumidores if p.solicitudes_del_mes]
    if con_detalle:
        elementos += [
            HRFlowable(width="100%", thickness=1, color=azul),
            Spacer(1, 0.3 * cm),
            Paragraph("DETALLE DE SOLICITUDES DEL MES", subtitulo_style),
        ]
    for perfil in con_detalle:
        d = fila_consumidor(perfil)
        elementos.append(Paragraph(f"{d['nombre']} — {d['ci']}", bold_style))
        filas_sol = [["N° Solicitud", "Fecha", "Estado", "Combustible", "Lit. Sol.", "Lit. Des.", "Estación"]]
        for s in perfil.solicitudes_del_mes:
            filas_sol.append([
                str(s.id_publico)[:8].upper(),
                formatear_fecha(s.fecha_creacion),
                s.get_estado_display(),
                s.get_tipo_combustible_display(),
                f"{s.litros_solicitados} L",
                f"{s.litros_despachados or 0} L",
                s.estacion_servicio.nombre if s.estacion_servicio else "—",
            ])
        tabla_sol = Table(filas_sol, repeatRows=1)
        tabla_sol.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), azul2),
            ("TEXTCOLOR", (0, 0), (-1, 0), blanco),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [blanco, gris]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elementos += [tabla_sol, Spacer(1, 0.4 * cm)]

    doc.build(elementos)
    return buffer.getvalue()
