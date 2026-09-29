# apps/solicitudes/services/formato_reportes.py
#
# Formato común de los archivos de reportes (solicitudes y consumidores):
#
#   PDF   → A4 apaisado; encabezado con título, período y filtros;
#           recuadro de resumen; tablas con encabezado sombreado, filas
#           alternadas, columnas de ancho fijo y texto ajustado
#           (Paragraph); pie "Página X de Y" con la fecha de generación.
#   Excel → hoja "Resumen" (título + indicadores) y hojas de tabla con
#           encabezado en negrita y fijo, autofiltro, anchos ajustados,
#           fechas como fechas reales y litros como números.
#
# Colores de la paleta de la app (tokens de index.css).

import io
from datetime import datetime
from xml.sax.saxutils import escape

from django.utils import timezone

from core.fechas import ahora_local, formatear_fecha_hora


NAVBAR    = "#17212B"
FILA_ALT  = "#F1F5F9"
BORDE     = "#CBD5E1"
TEXTO_SEC = "#6B7280"
FONDO_RES = "#F8FAFC"


# ================================================
# PDF
# ================================================

def _estilos():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle

    return {
        "titulo":    ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=15,
                                    textColor=colors.white, leading=19),
        "enc_linea": ParagraphStyle("enc_linea", fontName="Helvetica", fontSize=9,
                                    textColor=colors.HexColor("#E2E8F0"), leading=12),
        "subtitulo": ParagraphStyle("subtitulo", fontName="Helvetica-Bold", fontSize=11,
                                    textColor=colors.HexColor(NAVBAR), spaceBefore=8, spaceAfter=5),
        "celda":     ParagraphStyle("celda", fontName="Helvetica", fontSize=7.5, leading=9.5),
        "celda_num": ParagraphStyle("celda_num", fontName="Helvetica", fontSize=7.5, leading=9.5,
                                    alignment=TA_RIGHT),
        "cabecera":  ParagraphStyle("cabecera", fontName="Helvetica-Bold", fontSize=7.5, leading=9.5,
                                    textColor=colors.white),
        "res_label": ParagraphStyle("res_label", fontName="Helvetica", fontSize=7.5, leading=9,
                                    textColor=colors.HexColor(TEXTO_SEC), alignment=TA_CENTER),
        "res_valor": ParagraphStyle("res_valor", fontName="Helvetica-Bold", fontSize=13, leading=16,
                                    textColor=colors.HexColor(NAVBAR), alignment=TA_CENTER),
        "nota":      ParagraphStyle("nota", fontName="Helvetica-Oblique", fontSize=7.5, leading=10,
                                    textColor=colors.HexColor(TEXTO_SEC), spaceBefore=4),
        "bloque":    ParagraphStyle("bloque", fontName="Helvetica-Bold", fontSize=9, leading=12,
                                    textColor=colors.HexColor(NAVBAR), spaceBefore=6, spaceAfter=3),
    }


def texto(valor) -> str:
    """Texto seguro para Paragraph (escapa &, <, >)."""
    return escape("" if valor is None else str(valor))


def subtitulo(txt: str):
    from reportlab.platypus import Paragraph
    return Paragraph(texto(txt), _estilos()["subtitulo"])


def nota(txt: str):
    from reportlab.platypus import Paragraph
    return Paragraph(texto(txt), _estilos()["nota"])


def encabezado_bloque(txt: str):
    """Encabezado de una subsección (ej. un consumidor en el detalle)."""
    from reportlab.platypus import Paragraph
    return Paragraph(texto(txt), _estilos()["bloque"])


def _encabezado(titulo: str, lineas: list[str]):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    est = _estilos()
    contenido = [Paragraph("AGENCIA NACIONAL DE HIDROCARBUROS — BOLIVIA", est["enc_linea"]),
                 Paragraph(texto(titulo), est["titulo"])]
    contenido += [Paragraph(texto(l), est["enc_linea"]) for l in lineas]
    t = Table([[c] for c in contenido], colWidths=["100%"])
    t.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), colors.HexColor(NAVBAR)),
        ("LEFTPADDING",   (0, 0), (-1, -1), 12),
        ("TOPPADDING",    (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, -1), (-1, -1), 10),
        ("TOPPADDING",    (0, 1), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -2), 1),
    ]))
    return t


def recuadro_resumen(indicadores: list[tuple[str, str]]):
    """Indicadores en una fila de celdas: valor grande arriba, rótulo abajo."""
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    est = _estilos()
    celdas = [
        [Paragraph(texto(valor), est["res_valor"]), Paragraph(texto(label), est["res_label"])]
        for label, valor in indicadores
    ]
    t = Table([celdas], colWidths=[f"{100 / len(indicadores)}%"] * len(indicadores))
    t.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), colors.HexColor(FONDO_RES)),
        ("BOX",           (0, 0), (-1, -1), 0.8, colors.HexColor(BORDE)),
        ("LINEAFTER",     (0, 0), (-2, -1), 0.5, colors.HexColor(BORDE)),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING",    (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return t


def tabla(cabeceras: list[str], filas: list[list], anchos_cm: list[float],
          numericas: frozenset[int] = frozenset()):
    """
    Tabla de ancho fijo con texto ajustado. `numericas`: índices de
    columnas alineadas a la derecha.
    """
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Table, TableStyle

    est = _estilos()
    datos = [[Paragraph(texto(c), est["cabecera"]) for c in cabeceras]]
    for fila in filas:
        datos.append([
            Paragraph(texto(v), est["celda_num"] if i in numericas else est["celda"])
            for i, v in enumerate(fila)
        ])
    t = Table(datos, colWidths=[a * cm for a in anchos_cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND",     (0, 0), (-1, 0), colors.HexColor(NAVBAR)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(FILA_ALT)]),
        ("GRID",           (0, 0), (-1, -1), 0.4, colors.HexColor(BORDE)),
        ("VALIGN",         (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING",     (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING",  (0, 0), (-1, -1), 3),
    ]))
    return t


def _canvas_numerado(generado: str):
    """Canvas que dibuja 'Página X de Y · Generado: …' en cada página."""
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas

    class CanvasNumerado(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                ancho, _ = self._pagesize
                self.setFont("Helvetica", 7.5)
                self.setFillColor(colors.HexColor(TEXTO_SEC))
                self.drawString(42, 22, f"Generado: {generado}")
                self.drawRightString(ancho - 42, 22, f"Página {self._pageNumber} de {total}")
                super().showPage()
            super().save()

    return CanvasNumerado


def construir_pdf(titulo: str, lineas_encabezado: list[str], elementos: list) -> bytes:
    """Arma el PDF apaisado con encabezado + elementos + pie numerado."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.3 * cm, bottomMargin=1.5 * cm,
        title=titulo, author="ANH Bolivia",
    )
    doc.build(
        [_encabezado(titulo, lineas_encabezado), Spacer(1, 0.35 * cm)] + elementos,
        canvasmaker=_canvas_numerado(formatear_fecha_hora(ahora_local())),
    )
    return buffer.getvalue()


# Ancho útil de la hoja apaisada (A4 29.7 cm − 2 × 1.5 cm de margen)
ANCHO_UTIL_CM = 26.7


# ================================================
# EXCEL
# ================================================

FMT_FECHA      = "DD/MM/YYYY"
FMT_FECHA_HORA = "DD/MM/YYYY HH:MM"
FMT_ENTERO     = "#,##0"


def valor_excel(valor):
    """Datetimes aware → hora local sin zona (Excel no guarda zonas)."""
    if isinstance(valor, datetime):
        return timezone.localtime(valor).replace(tzinfo=None) if timezone.is_aware(valor) else valor
    return valor


def nuevo_libro():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def hoja_resumen(wb, titulo: str, lineas: list[str], indicadores: list[tuple[str, object]]):
    """Hoja "Resumen": título, período/filtros y los indicadores (rótulo | valor)."""
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet("Resumen")
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14)
    fila = 2
    for linea in lineas + [f"Generado: {formatear_fecha_hora(ahora_local())}"]:
        ws.cell(row=fila, column=1, value=linea).font = Font(italic=True, color="6B7280")
        fila += 1
    fila += 1
    ws.cell(row=fila, column=1, value="Indicador").font = Font(bold=True)
    ws.cell(row=fila, column=2, value="Valor").font = Font(bold=True)
    for label, valor in indicadores:
        fila += 1
        ws.cell(row=fila, column=1, value=label)
        c = ws.cell(row=fila, column=2, value=valor)
        c.alignment = Alignment(horizontal="right")
        if isinstance(valor, int):
            c.number_format = FMT_ENTERO
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18
    return ws


def hoja_tabla(wb, nombre: str, cabeceras: list[str], filas: list[list],
               anchos: list[float], formatos: dict[int, str] | None = None):
    """
    Hoja con una tabla: encabezado en negrita, sombreado y fijo (freeze
    panes), autofiltro y formatos por columna (índice → number_format).
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    formatos = formatos or {}
    ws = wb.create_sheet(nombre)
    ws.append(cabeceras)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVBAR.lstrip("#"))
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for fila in filas:
        ws.append([valor_excel(v) for v in fila])

    for col, fmt in formatos.items():
        letra = get_column_letter(col + 1)
        for celda in ws[letra][1:]:
            celda.number_format = fmt

    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabeceras))}{max(ws.max_row, 1)}"
    return ws


def libro_a_bytes(wb) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

