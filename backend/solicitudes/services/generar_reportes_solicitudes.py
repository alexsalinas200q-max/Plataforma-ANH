# apps/solicitudes/services/generar_reportes_solicitudes.py

import io
from django.db.models import Sum

from core.fechas import ahora_local, formatear_fecha_hora


# ------------------------------------------------
# HELPER — QUERYSET FILTRADO
# ------------------------------------------------

def _get_solicitudes(filtros: dict):
    """
    Retorna queryset de Solicitud aplicando los filtros
    que vienen desde la vista (ReporteSolicitudesView).
    """
    from solicitudes.models import Solicitud

    qs = Solicitud.objects.select_related(
        "consumidor__user",
        "estacion_servicio",
        "municipio",
    ).order_by("-fecha_creacion")

    fecha_desde  = filtros.get("fecha_desde")
    fecha_hasta  = filtros.get("fecha_hasta")
    estado       = filtros.get("estado")
    combustible  = filtros.get("combustible")
    estacion_id  = filtros.get("estacion_id")

    if fecha_desde:
        qs = qs.filter(fecha_creacion__date__gte=fecha_desde)
    if fecha_hasta:
        qs = qs.filter(fecha_creacion__date__lte=fecha_hasta)
    if estado:
        qs = qs.filter(estado=estado)
    if combustible:
        qs = qs.filter(tipo_combustible=combustible)
    if estacion_id:
        qs = qs.filter(estacion_servicio_id=estacion_id)

    return qs


def _fila_solicitud(s):
    """
    Extrae los datos de una solicitud para una fila del reporte.
    """
    consumidor_nombre = "—"
    if hasattr(s, "consumidor") and s.consumidor:
        if hasattr(s.consumidor, "user") and s.consumidor.user:
            consumidor_nombre = s.consumidor.user.nombre_completo()
        elif hasattr(s.consumidor, "nombre_completo"):
            consumidor_nombre = s.consumidor.nombre_completo

    return {
        "id_publico":       str(s.id_publico)[:8].upper(),
        "consumidor":       consumidor_nombre,
        "combustible":      s.get_tipo_combustible_display(),
        "litros_sol":       s.litros_solicitados,
        "litros_apr":       s.litros_aprobados or 0,
        "litros_des":       s.litros_despachados or 0,
        "estado":           s.get_estado_display(),
        "estado_raw":       s.estado,
        "municipio":        s.municipio.nombre if s.municipio else "—",
        "estacion":         s.estacion_servicio.nombre if s.estacion_servicio else "—",
        "fecha_creacion":   formatear_fecha_hora(s.fecha_creacion),
        "fecha_aprobacion": formatear_fecha_hora(s.fecha_aprobacion),
        "fecha_despacho":   formatear_fecha_hora(s.fecha_despacho),
    }


def _descripcion_filtros(filtros: dict) -> str:
    """Genera un texto legible con los filtros aplicados."""
    partes = []
    if filtros.get("fecha_desde"):
        partes.append(f"Desde: {filtros['fecha_desde']}")
    if filtros.get("fecha_hasta"):
        partes.append(f"Hasta: {filtros['fecha_hasta']}")
    if filtros.get("estado"):
        partes.append(f"Estado: {filtros['estado']}")
    if filtros.get("combustible"):
        partes.append(f"Combustible: {filtros['combustible']}")
    if filtros.get("estacion_id"):
        partes.append(f"Estación ID: {filtros['estacion_id']}")
    return " | ".join(partes) if partes else "Sin filtros (todos)"


# ------------------------------------------------
# GENERAR EXCEL DE SOLICITUDES
# ------------------------------------------------

def generar_excel_solicitudes(filtros: dict) -> bytes:
    """
    Genera un archivo Excel con el reporte de solicitudes.
    Retorna bytes del archivo.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Solicitudes"

    # Estilos
    header_font    = Font(bold=True, color="FFFFFF", size=11)
    header_fill    = PatternFill("solid", fgColor="1a3a5c")
    subheader_fill = PatternFill("solid", fgColor="2d6a9f")
    center         = Alignment(horizontal="center", vertical="center")
    thin_border    = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )

    # Colores por estado
    estado_fills = {
        "PENDIENTE":  PatternFill("solid", fgColor="FFF3CD"),
        "OBSERVADA":  PatternFill("solid", fgColor="FFE0CC"),
        "APROBADA":   PatternFill("solid", fgColor="D4EDDA"),
        "DESPACHADA": PatternFill("solid", fgColor="CCE5FF"),
        "RECHAZADA":  PatternFill("solid", fgColor="F8D7DA"),
        "CANCELADA":  PatternFill("solid", fgColor="E2E3E5"),
        "EXPIRADA":   PatternFill("solid", fgColor="E2E3E5"),
    }

    # Título
    ws.merge_cells("A1:L1")
    ws["A1"] = (
        f"REPORTE DE SOLICITUDES ANH — "
        f"{_descripcion_filtros(filtros)} — "
        f"Generado: {formatear_fecha_hora(ahora_local())}"
    )
    ws["A1"].font      = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill      = header_fill
    ws["A1"].alignment = center

    # Cabeceras
    headers = [
        "N° Solicitud", "Consumidor", "Combustible",
        "Litros Sol.", "Litros Apr.", "Litros Des.",
        "Estado", "Municipio", "Estación",
        "Fecha Creación", "Fecha Aprobación", "Fecha Despacho",
    ]

    for col, header in enumerate(headers, 1):
        cell            = ws.cell(row=2, column=col, value=header)
        cell.font       = header_font
        cell.fill       = subheader_fill
        cell.alignment  = center
        cell.border     = thin_border

    # Datos
    solicitudes = _get_solicitudes(filtros)
    row = 3

    for s in solicitudes:
        datos = _fila_solicitud(s)
        valores = [
            datos["id_publico"],
            datos["consumidor"],
            datos["combustible"],
            datos["litros_sol"],
            datos["litros_apr"],
            datos["litros_des"],
            datos["estado"],
            datos["municipio"],
            datos["estacion"],
            datos["fecha_creacion"],
            datos["fecha_aprobacion"],
            datos["fecha_despacho"],
        ]

        for col, valor in enumerate(valores, 1):
            cell           = ws.cell(row=row, column=col, value=valor)
            cell.border    = thin_border
            cell.alignment = Alignment(vertical="center")

        # Colorear fila según estado
        fill = estado_fills.get(datos["estado_raw"])
        if fill:
            for col in range(1, len(valores) + 1):
                ws.cell(row=row, column=col).fill = fill

        row += 1

    # Fila de totales
    totales = solicitudes.aggregate(
        total_sol=Sum("litros_solicitados"),
        total_apr=Sum("litros_aprobados"),
        total_des=Sum("litros_despachados"),
    )

    ws.cell(row=row, column=1, value="TOTALES").font = Font(bold=True)
    ws.cell(row=row, column=2, value=f"{solicitudes.count()} solicitudes").font = Font(bold=True)
    ws.cell(row=row, column=4, value=totales["total_sol"] or 0).font = Font(bold=True)
    ws.cell(row=row, column=5, value=totales["total_apr"] or 0).font = Font(bold=True)
    ws.cell(row=row, column=6, value=totales["total_des"] or 0).font = Font(bold=True)

    for col in range(1, len(headers) + 1):
        ws.cell(row=row, column=col).border = thin_border

    # Ajustar anchos
    anchos = [14, 28, 14, 12, 12, 12, 14, 16, 20, 18, 18, 18]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ------------------------------------------------
# GENERAR PDF DE SOLICITUDES
# ------------------------------------------------

def generar_pdf_solicitudes(filtros: dict) -> bytes:
    """
    Genera un PDF con el reporte de solicitudes.
    Incluye tabla con todas las solicitudes filtradas y totales.
    Retorna bytes del archivo.
    """
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle,
        Paragraph, Spacer,
    )
    from reportlab.lib.enums import TA_CENTER

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    azul   = colors.HexColor("#1a3a5c")
    azul2  = colors.HexColor("#2d6a9f")
    gris   = colors.HexColor("#f2f2f2")
    blanco = colors.white

    estado_colors = {
        "PENDIENTE":  colors.HexColor("#FFF3CD"),
        "OBSERVADA":  colors.HexColor("#FFE0CC"),
        "APROBADA":   colors.HexColor("#D4EDDA"),
        "DESPACHADA": colors.HexColor("#CCE5FF"),
        "RECHAZADA":  colors.HexColor("#F8D7DA"),
        "CANCELADA":  colors.HexColor("#E2E3E5"),
        "EXPIRADA":   colors.HexColor("#E2E3E5"),
    }

    titulo_style = ParagraphStyle(
        "titulo", fontSize=14, textColor=blanco,
        alignment=TA_CENTER, fontName="Helvetica-Bold",
    )
    subtitulo_style = ParagraphStyle(
        "subtitulo", fontSize=10, textColor=azul,
        alignment=TA_CENTER, fontName="Helvetica-Bold",
        spaceAfter=6,
    )

    elementos = []

    # Encabezado
    encabezado = Table(
        [[Paragraph(
            f"AGENCIA NACIONAL DE HIDROCARBUROS — BOLIVIA<br/>"
            f"Reporte de Solicitudes<br/>"
            f"<font size=9>{_descripcion_filtros(filtros)} | "
            f"Generado: {formatear_fecha_hora(ahora_local())}</font>",
            titulo_style,
        )]],
        colWidths=["100%"],
    )
    encabezado.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), azul),
        ("TOPPADDING",    (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING",   (0, 0), (-1, -1), 10),
    ]))
    elementos.append(encabezado)
    elementos.append(Spacer(1, 0.4 * cm))

    # Datos
    solicitudes = _get_solicitudes(filtros)

    totales = solicitudes.aggregate(
        total_sol=Sum("litros_solicitados"),
        total_apr=Sum("litros_aprobados"),
        total_des=Sum("litros_despachados"),
    )

    elementos.append(Paragraph(
        f"Total: {solicitudes.count()} solicitudes | "
        f"Litros solicitados: {totales['total_sol'] or 0} L | "
        f"Aprobados: {totales['total_apr'] or 0} L | "
        f"Despachados: {totales['total_des'] or 0} L",
        subtitulo_style,
    ))
    elementos.append(Spacer(1, 0.2 * cm))

    # Tabla
    cabeceras = [
        "N° Sol.", "Consumidor", "Comb.",
        "Lit. Sol.", "Lit. Apr.", "Lit. Des.",
        "Estado", "Municipio", "Estación",
        "Fecha", "Aprobación",
    ]

    filas = [cabeceras]
    filas_datos = []

    for s in solicitudes:
        datos = _fila_solicitud(s)
        filas_datos.append(datos)
        filas.append([
            datos["id_publico"],
            datos["consumidor"],
            datos["combustible"],
            f"{datos['litros_sol']} L",
            f"{datos['litros_apr']} L",
            f"{datos['litros_des']} L",
            datos["estado"],
            datos["municipio"],
            datos["estacion"],
            datos["fecha_creacion"].split(" ")[0] if datos["fecha_creacion"] != "—" else "—",
            datos["fecha_aprobacion"].split(" ")[0] if datos["fecha_aprobacion"] != "—" else "—",
        ])

    tabla = Table(filas, repeatRows=1)
    tabla.setStyle(TableStyle([
        # Encabezado
        ("BACKGROUND",    (0, 0), (-1, 0), azul2),
        ("TEXTCOLOR",     (0, 0), (-1, 0), blanco),
        ("FONTNAME",      (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, 0), 7.5),
        ("ALIGN",         (0, 0), (-1, 0), "CENTER"),
        ("TOPPADDING",    (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        # Filas
        ("FONTNAME",      (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE",      (0, 1), (-1, -1), 7),
        ("ALIGN",         (3, 1), (5, -1),  "CENTER"),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING",    (0, 1), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 3),
        # Bordes
        ("GRID",          (0, 0), (-1, -1), 0.5, colors.grey),
        ("BOX",           (0, 0), (-1, -1), 1, azul),
    ]))

    # Colorear filas según estado
    for i, datos in enumerate(filas_datos, 1):
        color = estado_colors.get(datos["estado_raw"])
        if color:
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, i), (-1, i), color),
            ]))

    elementos.append(tabla)

    doc.build(elementos)
    buffer.seek(0)
    return buffer.getvalue()